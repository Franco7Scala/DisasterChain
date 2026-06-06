import argparse
import csv
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from support.constants import *


def _stringify(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return " | ".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _load_dataset(path):
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("The final dataset JSON must contain an object keyed by disaster id.")
    return data


def _event_metadata(record, event_number):
    news_data = record.get("news_data") or {}
    search_metadata = news_data.get("search_metadata") or {}
    articles = news_data.get("articles") or []

    return {
        "benchmark_event_number": event_number,
        "disaster_id": record.get("disaster_id", ""),
        "country": record.get("country", ""),
        "region": record.get("region", ""),
        "disaster_type": record.get("disaster_type", ""),
        "start_date": record.get("start_date", ""),
        "latitude": record.get("latitude", ""),
        "longitude": record.get("longitude", ""),
        "news_engine_version": search_metadata.get("news_engine_version", ""),
        "total_articles_for_event": len(articles),
        "sources_successfully_resolved": _stringify(
            search_metadata.get("sources_successfully_resolved", [])
        ),
    }


def _article_metadata(article, article_index):
    return {
        "article_index": article_index,
        "source": article.get("source", ""),
        "title": article.get("title", ""),
        "url": article.get("url", ""),
        "raw_text": article.get("raw_text", ""),
        "raw_text_status": article.get("raw_text_status", ""),
        "raw_text_length": article.get("raw_text_length", ""),
        "raw_text_url": article.get("raw_text_url", ""),
        "relevance_score": article.get("relevance_score", ""),
        "relevance_threshold": article.get("relevance_threshold", ""),
        "confidence": article.get("confidence", ""),
        "is_non_article_candidate": article.get("is_non_article_candidate", ""),
        "relevance_reasons": _stringify(article.get("relevance_reasons", [])),
        "relevance_penalties": _stringify(article.get("relevance_penalties", [])),
        "search_query": article.get("search_query", ""),
        "query_precision": article.get("query_precision", ""),
        "published_at": article.get("published_at", ""),
    }


def build_benchmark_rows(dataset, limit):
    rows = []
    for event_number, (_, record) in enumerate(dataset.items(), start=1):
        if event_number > limit:
            break

        base_row = _event_metadata(record, event_number)
        articles = (record.get("news_data") or {}).get("articles") or []

        if not articles:
            row = dict(base_row)
            row.update(_article_metadata({}, ""))
            row["manual_label"] = ""
            row["manual_notes"] = ""
            rows.append(row)
            continue

        for article_index, article in enumerate(articles, start=1):
            row = dict(base_row)
            row.update(_article_metadata(article, article_index))
            row["manual_label"] = ""
            row["manual_notes"] = ""
            rows.append(row)

    return rows


def write_csv(rows, output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def write_excel(rows, output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "News review"

    sheet.append(CSV_COLUMNS)
    for row in rows:
        sheet.append([row.get(column, "") for column in CSV_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    manual_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True)
    manual_header_font = Font(color="000000", bold=True)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    manual_columns = {"manual_label", "manual_notes"}
    for column_name in manual_columns:
        column_index = CSV_COLUMNS.index(column_name) + 1
        header_cell = sheet.cell(row=1, column=column_index)
        header_cell.fill = manual_fill
        header_cell.font = manual_header_font

    widths = {
        "benchmark_event_number": 12,
        "disaster_id": 18,
        "country": 16,
        "region": 14,
        "disaster_type": 16,
        "start_date": 13,
        "latitude": 12,
        "longitude": 12,
        "news_engine_version": 22,
        "total_articles_for_event": 12,
        "sources_successfully_resolved": 28,
        "article_index": 10,
        "source": 24,
        "title": 55,
        "url": 55,
        "raw_text": 70,
        "raw_text_status": 20,
        "raw_text_length": 14,
        "raw_text_url": 55,
        "relevance_score": 12,
        "relevance_threshold": 12,
        "confidence": 14,
        "is_non_article_candidate": 18,
        "relevance_reasons": 38,
        "relevance_penalties": 38,
        "search_query": 55,
        "query_precision": 14,
        "published_at": 24,
        "manual_label": 16,
        "manual_notes": 45,
    }

    for index, column_name in enumerate(CSV_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column_name, 18)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[CSV_COLUMNS.index("manual_label")].fill = manual_fill
        row[CSV_COLUMNS.index("manual_notes")].fill = manual_fill

    max_row = max(sheet.max_row, 2)
    label_column = get_column_letter(CSV_COLUMNS.index("manual_label") + 1)
    validation = DataValidation(
        type="list",
        formula1='"correct,wrong,uncertain"',
        allow_blank=True,
    )
    validation.error = "Use one of: correct, wrong, uncertain."
    validation.errorTitle = "Invalid manual label"
    sheet.add_data_validation(validation)
    validation.add(f"{label_column}2:{label_column}{max_row}")

    url_column = CSV_COLUMNS.index("url") + 1
    for row_number in range(2, sheet.max_row + 1):
        url_cell = sheet.cell(row=row_number, column=url_column)
        if isinstance(url_cell.value, str) and url_cell.value.startswith("http"):
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 85

    workbook.save(output_path)


def parse_args():
    parser = argparse.ArgumentParser(
        description="Create manual review files for disaster news retrieval."
    )
    parser.add_argument(
        "--input",
        default=FINAL_DATASET_OUTPUT_PATH,
        help="Path to final_environmental_causal_dataset.json.",
    )
    parser.add_argument(
        "--output",
        default=DEFAULT_OUTPUT_PATH,
        help="Output CSV path.",
    )
    parser.add_argument(
        "--excel-output",
        default=DEFAULT_EXCEL_OUTPUT_PATH,
        help="Output Excel path.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=100,
        help="Number of disaster events to include.",
    )
    parser.add_argument(
        "--no-excel",
        action="store_true",
        help="Only write the CSV file.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    dataset = _load_dataset(args.input)
    rows = build_benchmark_rows(dataset, args.limit)
    write_csv(rows, args.output)
    if not args.no_excel:
        write_excel(rows, args.excel_output)
    print(
        f"Benchmark CSV created: {args.output} "
        f"({args.limit} events, {len(rows)} review rows)."
    )
    if not args.no_excel:
        print(f"Benchmark Excel created: {args.excel_output}")


if __name__ == "__main__":
    main()
