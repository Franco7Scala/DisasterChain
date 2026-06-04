import argparse
import csv
import os
from collections import Counter, defaultdict

from openpyxl import load_workbook

from support.constants import *



def _normalize_label(value):
    return str(value or "").strip().lower()


def _to_float(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _load_rows(path):
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    rows = []

    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        label = _normalize_label(row.get("manual_label"))
        if not label:
            continue
        if label not in VALID_LABELS:
            raise ValueError(
                f"Invalid manual_label '{label}' for disaster "
                f"{row.get('disaster_id', '')}. Use correct, wrong, or uncertain."
            )
        row["manual_label"] = label
        row["relevance_score"] = _to_float(row.get("relevance_score"))
        row["relevance_threshold"] = _to_float(row.get("relevance_threshold"))
        rows.append(row)

    return rows


def _safe_rate(numerator, denominator):
    if denominator == 0:
        return 0.0
    return numerator / denominator


def _label_counts(rows):
    return Counter(row["manual_label"] for row in rows)


def _precision_summary(rows):
    counts = _label_counts(rows)
    correct = counts["correct"]
    wrong = counts["wrong"]
    uncertain = counts["uncertain"]
    reviewed = correct + wrong + uncertain
    strict_denominator = correct + wrong

    return {
        "reviewed": reviewed,
        "correct": correct,
        "wrong": wrong,
        "uncertain": uncertain,
        "strict_precision": _safe_rate(correct, strict_denominator),
        "lenient_precision": _safe_rate(correct + 0.5 * uncertain, reviewed),
    }


def _group_rows(rows, key):
    grouped = defaultdict(list)
    for row in rows:
        grouped[str(row.get(key) or "missing")].append(row)
    return grouped


def _source_family(source):
    source = str(source or "missing")
    if source.startswith("Google News"):
        return "Google News"
    if source.startswith("DuckDuckGo"):
        return "DuckDuckGo"
    return source


def _group_by_source_family(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[_source_family(row.get("source"))].append(row)
    return grouped


def _score_bucket(score):
    if score is None:
        return "missing"
    if score < 5:
        return "<5"
    if score < 7:
        return "5-6"
    if score < 9:
        return "7-8"
    return "9+"


def _group_by_score_bucket(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[_score_bucket(row.get("relevance_score"))].append(row)
    return grouped


def _format_summary_line(name, rows):
    summary = _precision_summary(rows)
    return (
        f"{name}: reviewed={summary['reviewed']}, "
        f"correct={summary['correct']}, wrong={summary['wrong']}, "
        f"uncertain={summary['uncertain']}, "
        f"strict_precision={summary['strict_precision']:.2%}, "
        f"lenient_precision={summary['lenient_precision']:.2%}"
    )


def _wrong_cases(rows):
    cases = [
        row for row in rows
        if row.get("manual_label") == "wrong"
    ]
    return sorted(
        cases,
        key=lambda row: (
            row.get("relevance_score") if row.get("relevance_score") is not None else -1
        ),
        reverse=True,
    )


def _uncertain_cases(rows):
    return [
        row for row in rows
        if row.get("manual_label") == "uncertain"
    ]


def build_report(rows):
    lines = []
    lines.append("NEWS BENCHMARK REPORT")
    lines.append("=" * 60)
    lines.append(_format_summary_line("Overall", rows))
    lines.append("")

    lines.append("By source family")
    lines.append("-" * 60)
    for source, source_rows in sorted(_group_by_source_family(rows).items()):
        lines.append(_format_summary_line(source, source_rows))
    lines.append("")

    lines.append("By relevance score bucket")
    lines.append("-" * 60)
    bucket_order = ["missing", "<5", "5-6", "7-8", "9+"]
    by_score = _group_by_score_bucket(rows)
    for bucket in bucket_order:
        if bucket in by_score:
            lines.append(_format_summary_line(bucket, by_score[bucket]))
    lines.append("")

    lines.append("By confidence")
    lines.append("-" * 60)
    confidence_order = ["high", "medium", "low", "missing"]
    by_confidence = _group_rows(rows, "confidence")
    for confidence in confidence_order:
        if confidence in by_confidence:
            lines.append(_format_summary_line(confidence, by_confidence[confidence]))
    lines.append("")

    lines.append("By non-article candidate flag")
    lines.append("-" * 60)
    for flag, flag_rows in sorted(_group_rows(rows, "is_non_article_candidate").items()):
        lines.append(_format_summary_line(flag, flag_rows))
    lines.append("")

    lines.append("By disaster type")
    lines.append("-" * 60)
    for disaster_type, type_rows in sorted(_group_rows(rows, "disaster_type").items()):
        lines.append(_format_summary_line(disaster_type, type_rows))
    lines.append("")

    wrong_cases = _wrong_cases(rows)
    lines.append("Highest-score wrong cases")
    lines.append("-" * 60)
    for row in wrong_cases[:15]:
        lines.append(
            f"score={row.get('relevance_score')} | "
            f"{row.get('source')} | {row.get('disaster_id')} | "
            f"{row.get('title')} | {row.get('url')}"
        )
    lines.append("")

    uncertain_cases = _uncertain_cases(rows)
    lines.append("Uncertain cases")
    lines.append("-" * 60)
    for row in uncertain_cases[:15]:
        lines.append(
            f"score={row.get('relevance_score')} | "
            f"{row.get('source')} | {row.get('disaster_id')} | "
            f"{row.get('title')} | notes={row.get('manual_notes') or ''}"
        )

    return "\n".join(lines)


def write_report(report, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(report)


def write_wrong_cases(rows, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    columns = [
        "manual_label",
        "manual_notes",
        "relevance_score",
        "relevance_threshold",
        "confidence",
        "is_non_article_candidate",
        "source",
        "disaster_id",
        "country",
        "disaster_type",
        "start_date",
        "title",
        "url",
        "relevance_reasons",
        "relevance_penalties",
        "search_query",
    ]
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in _wrong_cases(rows):
            writer.writerow({column: row.get(column, "") for column in columns})


def parse_args():
    parser = argparse.ArgumentParser(
        description="Analyze the manually reviewed news benchmark Excel file."
    )
    parser.add_argument(
        "--input",
        default=DEFAULT_INPUT_PATH,
        help="Path to the reviewed benchmark Excel file.",
    )
    parser.add_argument(
        "--report-output",
        default=DEFAULT_REPORT_PATH,
        help="Path for the text report.",
    )
    parser.add_argument(
        "--wrong-output",
        default=DEFAULT_WRONG_OUTPUT_PATH,
        help="Path for the wrong-cases CSV.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    rows = _load_rows(args.input)
    report = build_report(rows)
    write_report(report, args.report_output)
    write_wrong_cases(rows, args.wrong_output)
    print(report)
    print("")
    print(f"Report written to: {args.report_output}")
    print(f"Wrong cases CSV written to: {args.wrong_output}")


if __name__ == "__main__":
    main()
